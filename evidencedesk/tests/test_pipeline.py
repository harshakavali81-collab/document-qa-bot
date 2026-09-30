import io
import json
import pytest
from pathlib import Path
from src.core import Index,chunk_document,extract,evidence_answer,ABSTAIN
from src.answering import validate,generate

@pytest.fixture
def idx():
    return Index([c for p in Path('data/sample_documents').glob('*.txt') for c in chunk_document(p.name,p.read_bytes())])

def test_retrieves_leave(idx):
    hits=idx.search('How many annual leave days are available?')
    assert '18 annual leave days' in hits[0][0].text
    assert hits[0][0].source=='people_policy.txt'

def test_unknown_question(idx):
    assert evidence_answer(idx.search('What is the orbital velocity of Neptune?'))==ABSTAIN

def test_empty_question(idx):
    assert idx.search(' ')==[]

def test_persistence(idx,tmp_path):
    idx.save(tmp_path)
    loaded=Index.load(tmp_path)
    assert loaded.search('annual leave')[0][0]==idx.search('annual leave')[0][0]

def test_duplicate_documents():
    c=chunk_document('x.txt',b'Annual leave is 18 days.')
    assert len(Index(c+c).chunks)==len(c)

def test_overlap_and_tail():
    text=' '.join(f'word{i}' for i in range(301))
    chunks=chunk_document('long.txt',text.encode(),size=120,overlap=25)
    assert chunks[0].text.split()[-25:]==chunks[1].text.split()[:25]
    assert chunks[-1].text.endswith('word300')

def test_no_text():
    with pytest.raises(ValueError,match='No readable'):
        extract('empty.txt',b' ')

def test_docx_table():
    from docx import Document
    doc=Document();doc.add_paragraph('Hello document');doc.add_table(rows=1,cols=1).cell(0,0).text='Revenue 100'
    buf=io.BytesIO();doc.save(buf)
    parts=extract('x.docx',buf.getvalue())
    assert ('paragraph 1','Hello document') in parts
    assert ('table 1','Revenue 100') in parts

def test_pdf_page_citation():
    # A minimal valid PDF with a real extractable text stream.
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
    writer=PdfWriter();page=writer.add_blank_page(width=612,height=792)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 50 700 Td (Annual leave is 18 days.) Tj ET')
    page[NameObject('/Contents')]=writer._add_object(stream)
    buf=io.BytesIO();writer.write(buf)
    assert extract('x.pdf',buf.getvalue())[0][0]=='page 1'
    assert '18 days' in chunk_document('x.pdf',buf.getvalue())[0].text

@pytest.mark.parametrize('content',[
    '{"answer":"Wrong [9]","citations":[9]}',
    '{"answer":"No citation","citations":[1]}',
    '{"answer":"Wrong [1]","citations":[true]}',
    'not json',
])
def test_bad_citations(content):
    with pytest.raises(RuntimeError):
        validate(content,2)

def test_valid_citation():
    assert validate('{"answer":"18 days [1]","citations":[1]}',2)=='18 days [1]'

def test_abstain_without_network():
    assert generate('unknown',[],None,None)==ABSTAIN

def test_ui_sample_flow():
    from streamlit.testing.v1 import AppTest
    at=AppTest.from_file(Path('app.py').resolve()).run(timeout=30)
    assert not at.exception
    at.button[0].click().run(timeout=30)
    assert not at.exception
    at.chat_input[0].set_value('How many annual leave days are available?').run(timeout=30)
    assert not at.exception
    assert '18 annual leave days' in at.session_state['history'][0]['answer']
