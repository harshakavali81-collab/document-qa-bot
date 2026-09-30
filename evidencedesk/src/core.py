"""Local retrieval with safe JSON persistence; no pickle loading."""
from dataclasses import dataclass, asdict
from pathlib import Path
import hashlib
import io
import json
import re
import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer

ABSTAIN = "I couldn't find enough information in the uploaded documents."
MODEL = 'sentence-transformers/all-MiniLM-L6-v2'

@dataclass
class Chunk:
    id: str
    source: str
    location: str
    text: str

def extract(name, data):
    """Return (location, text) pairs. DOCX/TXT have sections, not page numbers."""
    if len(data) > 10 * 1024 * 1024:
        raise ValueError('Maximum file size is 10 MB.')
    suffix = Path(name).suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > 250:
            raise ValueError('Maximum PDF length is 250 pages.')
        parts = [(f'page {i+1}', page.extract_text() or '') for i, page in enumerate(reader.pages)]
    elif suffix == '.docx':
        from docx import Document
        import zipfile
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(x.file_size for x in z.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Expanded DOCX exceeds 50 MB.')
        doc = Document(io.BytesIO(data))
        parts = [(f'paragraph {i+1}', p.text) for i,p in enumerate(doc.paragraphs)]
        parts += [(f'table {i+1}', '\n'.join(' | '.join(c.text for c in row.cells) for row in t.rows)) for i,t in enumerate(doc.tables)]
    elif suffix == '.txt':
        parts = [(f'section {i+1}', t) for i,t in enumerate(data.decode('utf-8-sig').split('\n\n'))]
    else:
        raise ValueError('Supported formats: PDF, DOCX, UTF-8 TXT.')
    parts = [(loc, re.sub(r'[ \t]+',' ',text).strip()) for loc,text in parts if text.strip()]
    if not parts:
        raise ValueError('No readable text. Scanned PDFs require OCR before upload.')
    if sum(len(t) for _,t in parts) > 2_000_000:
        raise ValueError('Extracted text exceeds 2 million characters.')
    return parts

def chunk_document(name, data, size=120, overlap=25):
    if not 0 <= overlap < size:
        raise ValueError('Require 0 <= overlap < chunk size.')
    digest = hashlib.sha256(data).hexdigest()[:12]
    chunks = []
    for location, text in extract(name, data):
        words = text.split()
        for start in range(0, len(words), size-overlap):
            piece = ' '.join(words[start:start+size])
            cid = hashlib.sha256(f'{name}:{digest}:{location}:{start}'.encode()).hexdigest()[:16]
            chunks.append(Chunk(cid, Path(name).name, location, piece))
            if start+size >= len(words):
                break
    return chunks

class Index:
    def __init__(self, chunks, backend='lexical'):
        if not chunks:
            raise ValueError('Add at least one readable document.')
        if backend not in ('lexical','semantic'):
            raise ValueError('Unknown retrieval backend.')
        self.chunks = list({c.id:c for c in chunks}.values())
        if len(self.chunks) > 10000:
            raise ValueError('Maximum 10,000 chunks per index.')
        self.backend = backend
        texts = [c.text for c in self.chunks]
        if backend == 'semantic':
            from sentence_transformers import SentenceTransformer
            self.encoder = SentenceTransformer(MODEL)
            self.vectors = self.encoder.encode(texts, normalize_embeddings=True)
        else:
            self.encoder = TfidfVectorizer(stop_words='english', ngram_range=(1,2), sublinear_tf=True)
            self.vectors = self.encoder.fit_transform(texts)

    def search(self, question, k=4, threshold=None):
        if not question.strip():
            return []
        if len(question) > 2000:
            raise ValueError('Question exceeds 2,000 characters.')
        if threshold is None:
            threshold = 0.12 if self.backend == 'lexical' else 0.35
        if self.backend == 'semantic':
            q = self.encoder.encode([question], normalize_embeddings=True)
            scores = (self.vectors @ q[0]).ravel()
        else:
            scores = (self.vectors @ self.encoder.transform([question]).T).toarray().ravel()
        return [(self.chunks[int(i)], float(scores[i])) for i in np.argsort(-scores)[:k] if scores[i] >= threshold]

    def save(self, directory):
        path = Path(directory)
        path.mkdir(parents=True, exist_ok=True)
        (path/'index.json').write_text(json.dumps({'version':1,'backend':self.backend,'model':MODEL,'chunks':[asdict(c) for c in self.chunks]},ensure_ascii=False),encoding='utf-8')
        if self.backend == 'semantic':
            np.save(path/'vectors.npy',self.vectors,allow_pickle=False)

    @classmethod
    def load(cls, directory):
        path = Path(directory)
        payload = json.loads((path/'index.json').read_text(encoding='utf-8'))
        if payload['version'] != 1 or payload['model'] != MODEL:
            raise ValueError('Index format or model mismatch. Rebuild the index.')
        chunks = [Chunk(**c) for c in payload['chunks']]
        if payload['backend'] == 'lexical':
            return cls(chunks,'lexical')
        if payload['backend'] != 'semantic':
            raise ValueError('Unknown retrieval backend.')
        from sentence_transformers import SentenceTransformer
        obj = cls.__new__(cls)
        obj.backend, obj.chunks = 'semantic',chunks
        obj.encoder = SentenceTransformer(MODEL)
        obj.vectors = np.load(path/'vectors.npy',allow_pickle=False)
        if obj.vectors.shape != (len(chunks),obj.encoder.get_sentence_embedding_dimension()):
            raise ValueError('Corrupt vector index. Rebuild it.')
        return obj

def evidence_answer(hits):
    if not hits:
        return ABSTAIN
    return 'Relevant document excerpts (not an AI-generated answer):\n\n' + '\n\n'.join(f'[{i}] {c.text}' for i,(c,_) in enumerate(hits,1))
